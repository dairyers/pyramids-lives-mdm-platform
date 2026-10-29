"""Builds the upload template for the stock-planning layer: data/stock_planning_template.xlsx (and a header-only CSV).
Run:  python tools/make_ingest_template.py
The Data sheet is empty on purpose, so nothing in it can be mistaken for real rows. Dropdowns and number checks stop
the usual faults (a typo in the record type, a text date, a negative quantity) while the file is being filled in."""
import pathlib

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

OUT = pathlib.Path(__file__).resolve().parents[1] / "data"
RECORD_TYPES = ["purchase_request", "order", "mdm_setup", "other"]
LAST_ROW = 5000

# name, required, type, unit, rule, example
COLUMNS = [
    ("date", True, "date", "YYYY-MM-DD", "Any day in the month the demand was logged. Use a real date, not text.", "2026-03-15"),
    ("country", True, "text", "code", "Country code exactly as you use it everywhere, for example SG, MY, TH.", "SG"),
    ("product", True, "text", "code", "Product line or item code. One consistent code per product.", "A"),
    ("quantity", True, "number", "units", "Units requested or ordered. Zero or more. No units in the cell, just the number.", "120"),
    ("record_type", True, "list", "", "purchase_request, order, mdm_setup or other. mdm_setup rows are kept out of demand.", "order"),
    ("on_hand", False, "number", "units", "Stock on hand at the end of that month. Lets us find real stockouts.", "340"),
    ("stockout_flag", False, "0 or 1", "", "1 if the product was out of stock at any point that month, otherwise 0.", "0"),
    ("order_date", False, "date", "YYYY-MM-DD", "Date you placed a supplier order. Lets us estimate lead time.", "2026-02-01"),
    ("receipt_date", False, "date", "YYYY-MM-DD", "Date that supplier order arrived.", "2026-03-04"),
    ("received_qty", False, "number", "units", "Units that arrived on that receipt.", "500"),
    ("unit_cost", False, "number", "currency per unit", "What one unit costs you. Needed for any cost figure.", "42.50"),
    ("price", False, "number", "currency per unit", "What you sell one unit for.", "61.00"),
    ("salvage", False, "number", "currency per unit", "What a leftover unit is worth if sold off. Zero if nothing.", "10.00"),
    ("order_cost", False, "number", "currency per order", "Cost of placing one supplier order (admin, freight, handling).", "85.00"),
    ("pack_size", False, "number", "units", "Orders must be a multiple of this.", "24"),
    ("moq", False, "number", "units", "Smallest quantity the supplier will accept.", "100"),
]
EXAMPLE_ROWS = [
    ["2026-01-12", "SG", "A", 120, "order"],
    ["2026-01-20", "SG", "A", 40, "purchase_request"],
    ["2026-01-25", "MY", "C", 500, "order"],
]
NUMBER_COLS = {"quantity", "on_hand", "received_qty", "unit_cost", "price", "salvage", "order_cost", "pack_size", "moq"}
DATE_COLS = {"date", "order_date", "receipt_date"}


def build(path):
    names = [c[0] for c in COLUMNS]
    wb = Workbook()

    ws = wb.active
    ws.title = "Data"
    required_fill = PatternFill("solid", fgColor="1F3A5F")
    optional_fill = PatternFill("solid", fgColor="D9E2F0")
    for i, (name, required, kind, unit, rule, example) in enumerate(COLUMNS, start=1):
        cell = ws.cell(row=1, column=i, value=name)
        cell.font = Font(bold=True, color="FFFFFF" if required else "1F3A5F")
        cell.fill = required_fill if required else optional_fill
        cell.alignment = Alignment(horizontal="center")
        cell.comment = Comment(f"{'Required' if required else 'Optional'} ({kind}{', ' + unit if unit else ''}). {rule} "
                               f"Example: {example}", "Template")
        ws.column_dimensions[cell.column_letter].width = max(14, len(name) + 4)
    ws.freeze_panes = "A2"

    def rule(kind, **kw):
        dv = DataValidation(type=kind, allow_blank=True, showErrorMessage=True, errorStyle="stop", **kw)
        ws.add_data_validation(dv)
        return dv

    for i, name in enumerate(names, start=1):
        col = ws.cell(row=1, column=i).column_letter
        span = f"{col}2:{col}{LAST_ROW}"
        if name in DATE_COLS:
            dv = rule("date", operator="between", formula1="DATE(2000,1,1)", formula2="DATE(2100,12,31)",
                      errorTitle="Date needed", error="Enter a real date such as 2026-03-15.")
            for r in range(2, LAST_ROW + 1):
                ws.cell(row=r, column=i).number_format = "yyyy-mm-dd"
        elif name in NUMBER_COLS:
            dv = rule("decimal", operator="greaterThanOrEqual", formula1="0",
                      errorTitle="Number needed", error="Enter a number that is zero or more, with no units or text.")
        elif name == "stockout_flag":
            dv = rule("whole", operator="between", formula1="0", formula2="1",
                      errorTitle="0 or 1", error="Enter 1 if the product was out of stock that month, otherwise 0.")
        elif name == "record_type":
            dv = rule("list", formula1='"' + ",".join(RECORD_TYPES) + '"',
                      errorTitle="Pick from the list", error="Use purchase_request, order, mdm_setup or other.")
        else:
            continue
        dv.add(span)

    guide = wb.create_sheet("Columns")
    for j, head in enumerate(["Column", "Required?", "Type", "Unit", "What to enter", "Example"], start=1):
        c = guide.cell(row=1, column=j, value=head)
        c.font, c.fill = Font(bold=True, color="FFFFFF"), required_fill
    for r, (name, required, kind, unit, rule_text, example) in enumerate(COLUMNS, start=2):
        for j, v in enumerate([name, "Yes" if required else "No", kind, unit, rule_text, example], start=1):
            guide.cell(row=r, column=j, value=v).alignment = Alignment(wrap_text=True, vertical="top")
    for letter, width in zip("ABCDEF", (16, 11, 10, 20, 78, 14)):
        guide.column_dimensions[letter].width = width
    notes = ["", "Rules that keep the upload clean",
             "1. Keep the headers in row 1 of the Data sheet exactly as they are. Do not rename or reorder the first five.",
             "2. One row per request or order line. Do not add totals, blank spacer rows or notes between rows.",
             "3. Use one unit for quantity across the whole sheet (for example units, not a mix of units and boxes).",
             "4. Leave optional cells blank when you do not have the value. Do not type 0, n/a or - for unknown.",
             "5. A product with fewer than 12 months of rows will get a simpler forecast and a warning.",
             "6. The Example sheet is for reading only. It is not read when you upload."]
    for k, text in enumerate(notes, start=len(COLUMNS) + 3):
        guide.cell(row=k, column=1, value=text).font = Font(bold=bool(text) and k == len(COLUMNS) + 4)

    ex = wb.create_sheet("Example")
    ex.cell(row=1, column=1, value="Example only: not read on upload").font = Font(bold=True, color="AA0808")
    for j, name in enumerate(names[:5], start=1):
        c = ex.cell(row=2, column=j, value=name)
        c.font, c.fill = Font(bold=True, color="FFFFFF"), required_fill
    for r, row in enumerate(EXAMPLE_ROWS, start=3):
        for j, v in enumerate(row, start=1):
            ex.cell(row=r, column=j, value=v)
    for letter in "ABCDE":
        ex.column_dimensions[letter].width = 18

    wb.save(path)
    (path.with_suffix(".csv")).write_text(",".join(names) + "\n", encoding="utf-8")


if __name__ == "__main__":
    target = OUT / "stock_planning_template.xlsx"
    build(target)
    print(f"wrote {target} and {target.with_suffix('.csv').name}")
