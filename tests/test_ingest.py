"""Tests for engine_ingest.py: reading, mapping, validating and reconciling a user's own data, then planning from it."""
import io
import pathlib

import numpy as np
import pandas as pd
import pytest
from openpyxl import load_workbook

import engine_ingest as eg
import engine_inventory as ei
from fixtures import CONFIG

ROOT = pathlib.Path(__file__).resolve().parents[1]
NOW = pd.Timestamp("2026-10-29", tz="UTC")
CANON = CONFIG["analytics"]["inventory"]["ingest"]["required"]


def synthetic(months=30, countries=("ZZ", "QQ"), products=("X1", "Y2"), seed=3):
    """A clean upload with countries and products the app has never seen, one yearly wave, mild growth and noise."""
    rng = np.random.default_rng(seed)
    rows = []
    for ci, c in enumerate(countries):
        for pi, p in enumerate(products):
            for m in range(months):
                base = 100 * (1 + ci + pi) * (1 + 0.01 * m) * (1 + 0.15 * np.sin(2 * np.pi * m / 12))
                date = pd.Timestamp("2024-01-01") + pd.DateOffset(months=m, days=int(rng.integers(0, 25)))
                rows.append({"date": date.date().isoformat(), "country": c, "product": p,
                             "quantity": round(float(base * rng.normal(1, 0.06)), 1), "record_type": "order"})
    return pd.DataFrame(rows)


def to_upload(df):
    out = df.copy()
    out["_row"] = np.arange(len(out)) + 2
    return out


def prepared(df, **kw):
    up = to_upload(df)
    mapping = eg.suggest_mapping([c for c in up.columns if c != "_row"], CONFIG)
    values = eg.suggest_record_types(up[mapping["record_type"]].dropna().unique(), CONFIG)
    return eg.prepare(up, mapping, values, CONFIG, **kw), mapping, values


# ----------------------------------------------------------------------------- reading and mapping
def test_template_headers_map_to_themselves():
    cols = list(pd.read_csv(ROOT / "data" / "stock_planning_template.csv").columns)
    m = eg.suggest_mapping(cols, CONFIG)
    assert all(m[c] == c for c in cols) and len(cols) == 20


def test_the_apps_own_sample_headers_are_recognised():
    m = eg.suggest_mapping(list(pd.read_csv(ROOT / "data" / "regional_orders.csv", nrows=2).columns), CONFIG)
    assert (m["date"], m["country"], m["product"], m["quantity"], m["record_type"]) == ("DocDate", "U_Country", "U_ProductLine", "Quantity", "RecordType")


def test_record_type_synonyms_and_unmapped_values():
    got = eg.suggest_record_types(["Order", "CONFIRMED_PO", "draft_po", "MDM setup", "enquiry", "banana"], CONFIG)
    assert got == {"Order": "order", "CONFIRMED_PO": "order", "draft_po": "purchase_request", "MDM setup": "mdm_setup",
                   "enquiry": "other", "banana": None}


def test_excel_template_round_trip_reads_the_data_sheet_not_the_example():
    wb = load_workbook(ROOT / "data" / "stock_planning_template.xlsx")
    ws = wb["Data"]
    for i, row in enumerate([["2026-01-12", "SG", "A", 120, "order"], ["2026-01-20", "SG", "A", 40, "purchase_request"]], start=2):
        for j, v in enumerate(row, start=1):
            ws.cell(row=i, column=j, value=v)
    buf = io.BytesIO()
    wb.save(buf)
    df = eg.read_upload(buf.getvalue(), "mine.xlsx")
    assert len(df) == 2 and list(df["_row"]) == [2, 3] and "record_type" in df.columns


def test_empty_template_has_no_rows_and_wrong_file_type_is_refused():
    assert len(eg.read_upload((ROOT / "data" / "stock_planning_template.xlsx").read_bytes(), "t.xlsx")) == 0
    with pytest.raises(ValueError):
        eg.read_upload(b"x", "notes.txt")


# ----------------------------------------------------------------------------- validation and reconciliation
def test_required_columns_must_be_mapped():
    up = to_upload(synthetic(14))
    mapping = eg.suggest_mapping(list(up.columns), CONFIG)
    mapping["quantity"] = None
    res = eg.prepare(up, mapping, {"order": "order"}, CONFIG)
    assert res["blocking"] and "quantity" in res["blocking"][0]


def test_every_bad_row_is_kept_with_a_reason_and_the_reconciliation_balances():
    df = synthetic(14).astype({"quantity": object})
    df.loc[0, "date"] = "not a date"
    df.loc[1, "quantity"] = -5
    df.loc[2, "quantity"] = "many"
    df.loc[3, "country"] = None
    df.loc[4, "record_type"] = "banana"
    df.loc[5, "record_type"] = "mdm_setup"
    df.loc[6, "record_type"] = "purchase_request"
    res, _, _ = prepared(df)
    t = res["table"]
    assert len(t) == len(df) and t["reason"].notna().sum() == 5
    assert set(t.loc[t["reason"].notna(), "reason"]) == {eg.REJECT["date"], eg.REJECT["quantity"], eg.REJECT["keys"], eg.REJECT["type"]}
    recon = eg.reconcile(t, "orders", CONFIG)
    assert int(recon["rows"].sum()) == len(df)
    assert recon.loc[recon["category"] == "mdm_setup", "treatment"].iloc[0].startswith("Excluded: waiting on master-data setup")
    assert recon.loc[recon["category"] == "purchase_request", "treatment"].iloc[0] == "Excluded: not in the chosen demand basis"


def test_setup_rows_never_enter_demand_under_either_basis():
    df = synthetic(14)
    df.loc[:9, "record_type"] = "mdm_setup"
    res, _, _ = prepared(df)
    for basis in ("orders", "requests"):
        frame = eg.demand_frame(res["table"], basis, CONFIG)
        assert len(frame) == (len(df) - 10 if basis == "orders" else 0)


def test_duplicates_are_flagged_and_only_dropped_when_asked():
    df = pd.concat([synthetic(14), synthetic(14).iloc[:3]], ignore_index=True)
    kept, _, _ = prepared(df)
    assert any(i[1] == "Duplicates" for i in kept["issues"]) and kept["table"]["reason"].isna().all()
    dropped, _, _ = prepared(df, drop_duplicates=True)
    assert (dropped["table"]["reason"] == eg.REJECT["duplicate"]).sum() == 3
    assert int(eg.reconcile(dropped["table"], "orders", CONFIG)["rows"].sum()) == len(df)


def test_gaps_short_history_outliers_and_mixed_units_are_reported_not_repaired():
    df = synthetic(30)
    df = df[~((df["country"] == "ZZ") & (df["product"] == "X1") & df["date"].str.startswith("2024-06"))]   # a missing month
    df = pd.concat([df, pd.DataFrame([{"date": "2025-03-05", "country": "ZZ", "product": "Y2", "quantity": 900000.0, "record_type": "order"},
                                      {"date": "2025-04-02", "country": "QQ", "product": "NEW", "quantity": 10.0, "record_type": "order"}])])
    res, _, _ = prepared(df.reset_index(drop=True))
    checks = {i[1] for i in res["issues"]}
    assert {"Missing months", "Short history", "Outliers"} <= checks
    assert len(eg.demand_frame(res["table"], "orders", CONFIG)) == len(df)                                      # nothing removed


def test_optional_columns_unused_for_now_are_said_so():
    df = synthetic(14)
    df["received_qty"] = 50.0
    res, _, _ = prepared(df)
    assert any(i[1] == "Column received_qty" and "not used by the plan yet" in i[2] for i in res["issues"])
    df["on_hand"] = 50.0
    assert not any(i[1] == "Column on_hand" for i in prepared(df)[0]["issues"])                               # stock is used now, so it is not reported as unused


def test_same_file_same_mapping_gives_the_same_result():
    a, _, _ = prepared(synthetic(20))
    b, _, _ = prepared(synthetic(20))
    pd.testing.assert_frame_equal(a["table"], b["table"])
    assert eg.data_id(a["table"], "orders") == eg.data_id(b["table"], "orders")


# ----------------------------------------------------------------------------- planning from an upload
def plan_from(df, **kw):
    res, _, _ = prepared(df)
    basis = eg.available_bases(res["table"], CONFIG)[0]
    clean = eg.demand_frame(res["table"], basis, CONFIG)
    recon = eg.reconcile(res["table"], basis, CONFIG)
    return ei.build_plan(None, CONFIG, basis=basis, demand=(clean, recon), data_label="Your data: test.csv", data_id=eg.data_id(res["table"], basis),
                         lead_days=res["lead_days"], now=NOW, **kw)


def test_unknown_countries_and_products_are_planned_and_labelled():
    p = plan_from(synthetic(30))
    t = p["table"]
    assert set(zip(t["country"], t["product"])) == {("ZZ", "X1"), ("ZZ", "Y2"), ("QQ", "X1"), ("QQ", "Y2")}
    assert (t["reorder_point"] > 0).all() and p["diagnostics"]["data_source"] == "Your data: test.csv"
    assert "file you uploaded" in p["diagnostics"]["warning"] and p["recon"]["rows"].sum() == 4 * 30
    assert t["order_quantity"].isna().all()                                                                    # no cost in the file: units only


def test_costs_and_lead_times_in_the_file_become_estimates_not_defaults():
    df = synthetic(30)
    df["unit_cost"], df["price"], df["order_cost"] = 20.0, 30.0, 60.0
    df["order_date"], df["receipt_date"] = "2026-01-01", "2026-02-15"
    p = plan_from(df)
    log = p["assumptions"].set_index(["country", "product", "assumption"])
    assert log.loc[("ZZ", "X1", "unit_cost"), "source"] == "estimated" and log.loc[("ZZ", "X1", "unit_cost"), "value"] == 20.0
    assert log.loc[("ZZ", "X1", "order_cost"), "value"] == 60.0 and log.loc[("ZZ", "X1", "order_cost"), "source"] == "estimated"
    assert log.loc[("ZZ", "X1", "lead_time_months"), "value"] == pytest.approx(45 / 30.4375)
    assert log.loc[("ZZ", "X1", "lead_time_months"), "source"] == "estimated" and log.loc[("ZZ", "X1", "holding_rate"), "source"] == "default"
    assert p["table"]["order_quantity"].notna().all()


def test_a_short_upload_is_planned_with_flags_not_refused():
    p = plan_from(synthetic(8))
    assert len(p["table"]) == 4 and p["table"]["flags"].str.contains("Under 12 months").all()


def test_same_upload_same_plan_and_a_different_upload_does_not_reuse_the_cache():
    cache = {}
    a = plan_from(synthetic(30), caches=cache)
    b = plan_from(synthetic(30), caches=cache)
    pd.testing.assert_frame_equal(a["table"], b["table"])
    c = plan_from(synthetic(30, seed=9), caches=cache)
    assert not np.allclose(a["table"]["forecast_next_12m"], c["table"]["forecast_next_12m"])


def test_the_apps_own_sample_file_can_be_uploaded_too():
    raw = pd.read_csv(ROOT / "data" / "regional_orders.csv")
    up = to_upload(raw)
    mapping = eg.suggest_mapping([c for c in up.columns if c != "_row"], CONFIG)
    values = eg.suggest_record_types(raw["RecordType"].unique(), CONFIG)
    res = eg.prepare(up, mapping, values, CONFIG)
    assert not res["blocking"] and int(eg.reconcile(res["table"], "orders", CONFIG)["rows"].sum()) == len(raw)


# ----------------------------------------------------------------------------- current stock, stockout flags and price breaks from a file
def test_the_latest_stock_figures_per_product_are_taken_from_the_file():
    df = synthetic(14)
    df["on_hand"] = np.arange(len(df), dtype=float)
    df["on_order"] = 5.0
    df.loc[df.index[-1], "on_hand"] = np.nan                                                              # the very last row has no stock figure
    res, _, _ = prepared(df)
    cur = eg.current_stock(res["table"])
    assert len(cur) == 4 and set(cur.columns) == {"country", "product", "on_hand", "on_order"} and (cur["on_order"] == 5.0).all()
    last_row = res["table"].sort_values(["date_parsed", "row"]).groupby(["country", "product"]).last()
    assert cur.set_index(["country", "product"])["on_hand"].notna().all() and last_row["on_hand"].notna().sum() >= 3
    assert eg.current_stock(prepared(synthetic(14))[0]["table"]).empty                                      # no stock column: nothing invented


def test_stockout_flags_in_the_file_are_listed():
    df = synthetic(14)
    df["stockout_flag"] = 0
    df.loc[[3, 40], "stockout_flag"] = 1
    flags = eg.stockout_flags(prepared(df)[0]["table"])
    assert len(flags) == 2 and all(len(f) == 3 for f in flags)


def test_price_breaks_are_read_from_the_template_sheet():
    wb = load_workbook(ROOT / "data" / "stock_planning_template.xlsx")
    ws = wb["Price breaks"]
    for i, row in enumerate([["ALL", "A", 300, 18.0], ["ALL", "A", 600, 16.0], ["SG", "A", 100, 19.0], ["ALL", "B", "x", 5.0], ["ALL", "B", 50, -1]], start=2):
        for j, v in enumerate(row, start=1):
            ws.cell(row=i, column=j, value=v)
    buf = io.BytesIO()
    wb.save(buf)
    got = eg.read_price_breaks(buf.getvalue(), "mine.xlsx")
    assert got == {("ALL", "A"): [(300.0, 18.0), (600.0, 16.0)], ("SG", "A"): [(100.0, 19.0)]}                 # bad rows ignored
    assert eg.read_price_breaks(b"x", "mine.csv") == {} and eg.read_price_breaks((ROOT / "data" / "stock_planning_template.xlsx").read_bytes(), "t.xlsx") == {}